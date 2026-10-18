"""放行与归档（规格 6.2 / 6.3 / 6.4 / 7.2）。

本任务钉的是**顺序**与**谓词**，两样都是「写对了也不报错、写错了静默卡死」
的东西：

- 闸门顺序必须是 ①④⑤②③⑥。④ 排在②前面是硬要求——重跑的行同时满足
  ②和④，②在前就把它挡死，而④永远走不到。
- ②的判据必须是 ``workflow_archived = 0``。写成 ``workflow_row < cursor``
  会把所有历史行算成没归档，整批永不放行，且只在用过重跑之后才发作。
- 归档要等打包结束（终态），不是等进程退出。
- **归档失败时「不置位」是可接受的，「置位」不可接受**：置位之后待归档行的取数
  （带 ``AND workflow_archived = 0``）再也查不到这一行，用户开着 Excel 跑完整批，
  结果表里对应行是空的，而且不报错。卡住能告诉用户，丢数据不能。
- tick **不许等一次完整的 write_back**：openpyxl 不支持增量写，宽表单次 2.2 秒，
  而 POLL_INTERVAL 只有 1 秒。归档跑在后台线程上。

这些用例直接调 ``_workflow_run_pass`` / ``_workflow_archive_pass``，
不真起 agent 进程——放行成败看的是库里那一行状态变了没。
"""
from __future__ import annotations

import threading
import time
from pathlib import Path

import pytest

import db
import dispatcher
import store


class _FakeProc:
    """占位的 AgentProc：只够 _reap 走一圈（既没中止、也没退出）。"""
    aborted = False
    exit_code = None

    def __init__(self, task_id):
        self.task_id = task_id
        self.proc = self

    def poll(self):
        return None


@pytest.fixture()
def disp(tmp_path, monkeypatch):
    """一个不真的 spawn 的调度器：把 _spawn 换成记账。

    用真的 ``Dispatcher(...)`` 而不是 ``__new__``：``running`` 是只读 property，
    而且它底下要 ``self._lock``——``__new__`` 出来的对象两者都没有，一碰就
    AttributeError，测的就不是闸门而是构造技巧了。
    """
    monkeypatch.setenv("PROTFORGE_WEB_NODE", "node")
    d = dispatcher.Dispatcher(db, max_concurrency=2, node_exe="node",
                              gpu_probe=lambda: 2)
    spawned: list[int] = []

    def fake_spawn(task_id, brief, work_dir):
        spawned.append(int(task_id))
        d._running[int(task_id)] = _FakeProc(task_id)
        return _FakeProc(task_id)

    monkeypatch.setattr(d, "_spawn", fake_spawn)
    monkeypatch.setattr(d, "_mark_failed",
                        lambda task_id, msg: store.set_status(
                            task_id, "failed", actor="system", message=msg))
    d._spawned = spawned
    yield d
    d._archive_stop()


def _user(name="u1", role="user"):
    c = db.connect()
    try:
        with c:
            c.execute("INSERT INTO users (username, password_hash, salt, role,"
                      " disabled, created_at) VALUES (?, 'h','s', ?, 0, 't')",
                      (name, role))
        return c.execute("SELECT id FROM users WHERE username=?", (name,)).fetchone()[0]
    finally:
        c.close()


def _workflow(owner_id, tmp_path, rows=5, name="wf", wid=9001, sheet="S1"):
    """建一个带真实 xlsx 源表的工作流，所以放行时能真的渲染出 brief。"""
    import workflow as wfmod
    root = Path(tmp_path) / "data" / "workflows"
    st = wfmod.staging_path(root)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["编号", "序列", "项目"])
    for i in range(rows):
        ws.append([f"ID{i}", f"SEQ{i}", "XBR"])
    wb.save(st)
    source = wfmod.finalize_upload(root, st, wid)
    wf = store.create_workflow(owner_id, name, "查 {{seq}}（项目 {{proj}}）",
                               sheet=sheet, source_path=str(source),
                               header_row=1, key_col=0, total_rows=rows,
                               writeback=1)
    store.set_workflow_columns(wf["id"], [
        {"col_index": 0, "header": "编号", "role": "key"},
        {"col_index": 1, "header": "序列", "role": "input", "slot": "seq"},
        {"col_index": 2, "header": "项目", "role": "context", "slot": "proj"},
        {"col_index": 3, "header": "状态", "role": "output", "field": "status"},
        {"col_index": 4, "header": "说明", "role": "output", "field": "summary"},
    ])
    store.update_workflow(wf["id"], status="running")
    return wf


def _first_row(wf, index=0):
    return store.list_workflow_rows(wf["id"], index)[0]


def _archived(task_id) -> int:
    return int(store.get_task(task_id)["workflow_archived"])


# --- 放行基本流 -------------------------------------------------------------

def test_admits_first_row_and_advances_cursor(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path)
    got = disp._workflow_run_pass()
    assert len(got) == 1
    assert store.get_workflow(wf["id"])["cursor"] == 1
    row = _first_row(wf)
    assert row["status"] == "running"          # claim_queued 把它推成 running
    assert "SEQ0" in row["brief"]


def test_row_task_owner_is_the_workflow_owner(disp, tmp_path):
    """闸门⑥ 判的是 owner 的队列，不是别人的。"""
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    row = _first_row(wf)
    assert row["owner_id"] == uid


def test_row_task_created_event_names_the_workflow(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path, name="抗体筛选")
    disp._workflow_run_pass()
    row = _first_row(wf)
    evs = store.list_events(row["id"])
    created = [e for e in evs if e["kind"] == "created"]
    assert len(created) == 1
    assert created[0]["actor"] == "system"
    assert "抗体筛选" in created[0]["message"] and "1" in created[0]["message"]


def test_brief_of_row_n_excludes_other_rows(disp, tmp_path):
    """§6.9 上下文独立：**第 N 行的 brief 里不含第 N-1 行的值**。

    这是「任务之间上下文相互独立」那句话唯一能变成测试的形状。天然成立
    （每行独立 task、独立 work 目录），但必须钉死——批量场景里一行漏进
    上一行的序列，整个后续批次的结果都是错的，而且**不报错**。
    """
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    disp._workflow_run_pass()
    r1 = _first_row(wf, 1)
    assert "SEQ1" in r1["brief"]
    assert "SEQ0" not in r1["brief"], "第 2 行的 brief 里混进了第 1 行的序列"


def test_serial_second_row_waits(disp, tmp_path):
    uid = _user()
    _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    assert disp._workflow_run_pass() == []      # 同一 tick 不放第二行


def test_gate_5_holds_even_when_the_row_is_already_archived(disp, tmp_path):
    """⑤ 与② 会互相遮蔽：未归档的行② 就挡住了，**看不出⑤ 在不在**。

    能把二者区分开的只有「已归档但还没进终态」的行——闸门⑤ 故意不过滤
    workflow_archived 正是为它准备的（加了过滤，已归档非终态的行会既过⑤又不过
    ②③，同一个工作流两行同时推进）。少了这条围栏，⑤ 被整个删掉、或被加上
    `workflow_archived = 0` 过滤，上面那些用例都照样全绿。
    """
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="running", workflow_archived=1)
    assert disp._workflow_run_pass() == [], "闸门⑤ 被② 遮蔽掉了"
    assert store.list_workflow_rows(wf["id"], 1) == []


def test_gate_1_is_rechecked_on_the_workflow_dict_itself(disp, tmp_path):
    """闸门① 有**两道**：取数（list_running_workflows 只给 running）与
    :meth:`_workflow_admit_one` 里的再判一次。

    后者是纵深防御，也正因为冗余，只有直接调内部方法才钉得住它——从 pass 那一层
    看，两道少任何一道结果都一样，全绿的测试证明不了它还在。
    """
    uid = _user()
    wf = _workflow(uid, tmp_path)
    store.update_workflow(wf["id"], status="paused")
    assert disp._workflow_admit_one(store.get_workflow(wf["id"])) is None
    assert disp._spawned == []


def test_next_row_after_previous_is_archived(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done")
    assert disp._workflow_run_pass() == [], "归档前不许放下一行"
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1
    got = disp._workflow_run_pass()
    assert len(got) == 1
    assert store.list_workflow_rows(wf["id"], 1)


def test_packaging_blocks_archiving(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="packaging")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 0
    store.update_task(r0["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1


# --- 闸门顺序（规格 6.4）----------------------------------------------------

def test_gate_4_precedes_gate_2(disp, tmp_path):
    """重跑的行同时满足②和④。④不排在前面 → 这条永远停在「已排队」。

    这条用例单独存在的理由：整理代码顺序的人不会想到去跑重跑用例。
    """
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    store.update_task(r0["id"], status="queued", workflow_archived=0)   # 管理员重跑
    disp._running.clear()      # 模型上管理员重跑时，上一次那个进程早就退出了
    got = disp._workflow_run_pass()
    assert len(got) == 1, "重跑的行没有被认领——④ 被②挡住了"
    assert store.get_task(r0["id"])["status"] == "running"
    assert store.get_workflow(wf["id"])["cursor"] == 1, "重跑不该再推游标"


def test_regular_tasks_go_first(disp, tmp_path):
    """顺序层让路：常规任务和工作流行任务同时有空槽时，先给常规的。"""
    uid = _user()
    wf = _workflow(uid, tmp_path)
    store.create_task("常规", "b", owner_id=uid)
    admitted = disp.tick()
    assert store.list_workflow_rows(wf["id"], 0) == []
    assert len(admitted) == 1


def test_gate_6_waits_for_owner_regular_tasks(disp, tmp_path):
    uid = _user()
    _workflow(uid, tmp_path)
    store.create_task("A 的常规任务", "b", owner_id=uid)
    assert disp._workflow_run_pass() == []


def test_gate_6_ignores_other_users_regular_tasks(disp, tmp_path):
    """owner_id 分错的话，A 的工作流会被 B 的任务挡住。"""
    a, b = _user("a"), _user("b")
    _workflow(a, tmp_path)
    store.create_task("B 的常规任务", "b", owner_id=b)
    assert len(disp._workflow_run_pass()) == 1


def test_gate_6_ignores_the_same_users_other_workflow_rows(disp, tmp_path):
    """闸门⑥ 只看**常规**任务。同一个人另一个工作流的行任务不该挡住这个。"""
    a = _user("a")
    wf_a = _workflow(a, tmp_path, name="A", wid=9101)
    wf_b = _workflow(a, tmp_path, name="B", wid=9102)
    other = store.create_task("B 的第 0 行", "b", owner_id=a,
                              workflow_id=wf_b["id"], workflow_row=0)
    store.update_task(other["id"], status="running")   # B 正在跑，不进 _running
    assert len(disp._workflow_run_pass()) == 1, "同用户的另一个工作流行任务挡住了闸门⑥"
    assert disp._spawned == [_first_row(wf_a)["id"]]
    assert store.get_workflow(wf_b["id"])["cursor"] == 0, "B 自己那一行是闸门⑤ 的事"


def test_concurrency_limit_is_enforced(disp, tmp_path):
    """_spawn 本身没有限流检查，全靠这条闸门。"""
    a, b = _user("a"), _user("b")
    disp._running[999] = object()
    disp.limit = 1
    _workflow(a, tmp_path, name="A", wid=9201)
    _workflow(b, tmp_path, name="B", wid=9202)
    assert disp._workflow_run_pass() == []


def test_paused_workflow_admits_nothing(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path)
    store.update_workflow(wf["id"], status="paused")
    # 闸门① 的取数侧也单独钉一下：它与 _workflow_admit_one 里的再判是冗余的，
    # 只断言 pass 不放行的话，两道里少任何一道都看不出来。
    assert store.list_running_workflows() == [], "闸门① 的取数侧没滤掉 paused"
    assert disp._workflow_run_pass() == []


def test_finished_when_cursor_exhausted(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path, rows=1)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    disp._workflow_run_pass()
    assert store.get_workflow(wf["id"])["status"] == "finished"


# --- 归档（规格 7.2）--------------------------------------------------------

def test_archive_marks_every_row_individually(disp, tmp_path):
    """一批两行都终态 → 两行都要置位，不许只处理第一行。"""
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    disp._workflow_run_pass()
    r1 = _first_row(wf, 1)
    store.update_task(r1["id"], status="done")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1
    assert _archived(r1["id"]) == 1


def test_locked_result_file_leaves_row_retryable(disp, tmp_path, monkeypatch):
    """文件被 Excel 占住 → **不置位**、留待重试。

    置位就等于宣告「这一行完事了」，而结果根本不在表里：待归档行的取数带
    ``AND workflow_archived = 0``，置位之后没有任何代码会再碰它，writeback_pending
    永远挂着没人清，用户看到的却是「跑完了」。
    """
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")

    monkeypatch.setattr(wfmod, "write_back",
                        lambda *a, **k: (_ for _ in ()).throw(
                            PermissionError(13, "被占用")))
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)

    assert _archived(r0["id"]) == 0, "被占时置了位——这一行将永远不会被重试"
    assert store.get_workflow(wf["id"])["writeback_pending"] == 1
    notes = [e for e in store.list_events(r0["id"]) if e["kind"] == "note"]
    assert notes, "用户看不到「关掉 Excel 后会自动重试」这条提示"
    assert "Excel" in notes[-1]["message"]


def test_locked_row_is_written_back_on_the_next_archive_pass(disp, tmp_path,
                                                            monkeypatch):
    """解锁后由**下一次归档顺手补写**——不是定时器重试。

    只测「被占时不置位」不够：那证明了它不丢，但也可能永远不补。少这一条的
    时候，实现可以是「永远标 pending、永远不写」而全部测试通过。
    """
    from openpyxl import load_workbook
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")

    real = wfmod.write_back
    state = {"locked": True}

    def flaky(*a, **k):
        if state["locked"]:
            raise PermissionError(13, "被占用")
        return real(*a, **k)

    monkeypatch.setattr(wfmod, "write_back", flaky)
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 0

    state["locked"] = False
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)

    assert _archived(r0["id"]) == 1
    assert store.get_workflow(wf["id"])["writeback_pending"] == 0
    result = Path(store.get_workflow(wf["id"])["result_path"])
    ws = load_workbook(result)["S1"]
    assert ws.cell(row=2, column=4).value == "done"
    assert ws.cell(row=2, column=5).value == "ok"


def test_source_shortened_marks_row_archived_with_error_event(disp, tmp_path):
    """源表被改短 → 重试一万次也等不到，必须置位，否则闸门② 整批死锁。"""
    from openpyxl import Workbook
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")
    # 把源表改成只剩表头：第 0 个数据行不存在了 → _sheet_row_number 抛 Invalid
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.append(["编号", "序列", "项目", "状态", "说明"])
    wb.save(store.get_workflow(wf["id"])["source_path"])

    disp._workflow_archive_pass()
    assert disp._archive_wait(30)

    assert _archived(r0["id"]) == 1, "数据损坏不置位的话闸门② 会永远等这一行"
    errs = [e for e in store.list_events(r0["id"]) if e["kind"] == "error"]
    assert errs, "失败原因没有落到用户可见的时间线上"
    assert store.workflow_has_unarchived(wf["id"]) is False, \
        "闸门② 的谓词还为真，整批就死在这一行上"


def test_gate_2_blocks_next_row_while_unarchived(disp, tmp_path, monkeypatch):
    """反向围栏：闸门② 在「有未归档的行」时确实不放行。

    钉的是**卡住是暂时的**：上一条钉住「被占时不置位」，这条钉住「不置位之后
    这一批真的停住、下一次归档成功之后又真的放行」。两条合起来才是完整契约；
    少任何一条，实现都可以是「永远不放行」或「永远放行」而全绿。
    """
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")

    real = wfmod.write_back
    state = {"locked": True}

    def flaky(*a, **k):
        if state["locked"]:
            raise PermissionError(13, "被占用")
        return real(*a, **k)

    monkeypatch.setattr(wfmod, "write_back", flaky)
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert disp._workflow_run_pass() == [], "结果还没落表就放下一行了"
    assert store.get_workflow(wf["id"])["cursor"] == 1

    state["locked"] = False
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1
    assert len(disp._workflow_run_pass()) == 1, "补写成功后闸门② 应当解除"


def test_archive_pass_exception_does_not_escape(disp, tmp_path):
    """一次异常不能越过 tick 的中线，否则所有常规任务停在排队中。"""
    uid = _user()
    _workflow(uid, tmp_path)
    disp._workflow_run_pass()

    def boom():
        raise RuntimeError("整批写挂了")
    disp._workflow_archive_pass = boom
    # tick 不该炸
    assert isinstance(disp.tick(), list)


def test_one_broken_row_does_not_stop_the_batch(disp, tmp_path, monkeypatch):
    """第 1 行的 write_back 抛异常，第 2 行照样归档，线程照样活着。"""
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")
    r1 = store.create_task("补的行", "b", owner_id=uid,
                           workflow_id=wf["id"], workflow_row=1)
    store.update_task(r1["id"], status="done", summary="ok2")

    real = wfmod.write_back

    def boom_first(*a, **k):
        if a[4][0][0] == 0:                 # jobs = [(row_index, col, value)]
            raise RuntimeError("openpyxl 崩了")
        return real(*a, **k)

    monkeypatch.setattr(wfmod, "write_back", boom_first)
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)

    assert _archived(r0["id"]) == 0
    assert _archived(r1["id"]) == 1, "一行失败把整批带走了"
    assert disp._archive_thread is not None and disp._archive_thread.is_alive(), \
        "归档线程被一个异常带走了"


# --- 归档不阻塞调度主路径 ----------------------------------------------------

def test_tick_does_not_wait_for_write_back(disp, tmp_path, monkeypatch):
    """POLL_INTERVAL 是 1 秒，而宽表单次 write_back 2.2 秒——tick 绝不能等它。

    这条不是性能断言而是**结构断言**：write_back 被事件门卡住不放行，tick 只要
    同步做归档就必然超时返回。
    """
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")

    entered = threading.Event()
    release = threading.Event()
    real = wfmod.write_back

    def slow(*a, **k):
        entered.set()
        if not release.wait(10):
            raise AssertionError("tick 在等一次完整的 write_back")
        return real(*a, **k)

    monkeypatch.setattr(wfmod, "write_back", slow)
    try:
        t0 = time.time()
        disp.tick()
        elapsed = time.time() - t0
        assert entered.wait(5), "归档线程根本没开工"
        assert elapsed < 3.0, f"tick 花了 {elapsed:.2f}s，归档还在主路径上"
        # 归档还没完成 → 这一行仍未归档 → 闸门② 仍不放行
        assert _archived(r0["id"]) == 0
        assert disp._workflow_run_pass() == []
    finally:
        release.set()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1


def test_row_in_flight_is_not_enqueued_twice(disp, tmp_path, monkeypatch):
    """竞态：tick 刚把行交给归档线程、还没置位就又扫到一次。

    ``workflow_archived = 0`` 这个谓词在这里**兜不住**：那一列要等整表写完才
    置位，中间有两秒。这两秒里每次 tick 都会把同一行重新捞出来——所以必须靠
    在途集合（或者别的等价手段）去重，否则同一个文件被同时写两次。
    """
    import workflow as wfmod
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")

    entered = threading.Event()
    release = threading.Event()
    real = wfmod.write_back
    calls = []

    def slow(*a, **k):
        calls.append(1)
        entered.set()
        release.wait(10)
        return real(*a, **k)

    monkeypatch.setattr(wfmod, "write_back", slow)
    try:
        first = disp._workflow_archive_pass()
        assert entered.wait(5)
        second = disp._workflow_archive_pass()
        third = disp._workflow_archive_pass()
        assert first == [r0["id"]]
        assert second == [] and third == [], "在途的行被重复排队了"
    finally:
        release.set()
    assert disp._archive_wait(30)
    assert len(calls) == 1, "同一个结果文件被同时写了两次"


# --- 常规队列排除工作流任务（规格 6.1）-------------------------------------

def test_queued_tasks_excludes_workflow_rows(disp, tmp_path):
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()                       # 行任务变成 running
    store.update_task(_first_row(wf)["id"], status="queued")   # 人工退回队列
    got = [int(r["id"]) for r in disp._queued_tasks()]
    assert _first_row(wf)["id"] not in got


# --- 回写的目标表一致性（规格 7.1）------------------------------------------

def _two_sheet_book(tmp_path):
    from openpyxl import Workbook
    src = Path(tmp_path) / "two.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.append(["编号", "状态"])
    ws.append(["ID0", ""])
    other = wb.create_sheet("S2")
    other.append(["编号", "状态"])
    other.append(["ID0", ""])
    wb.save(src)
    return src


def test_write_back_rejects_a_sheet_the_row_numbers_come_from_nowhere(tmp_path):
    """行号按第一张表算、值写进第二张表 = 静默写错行。必须报错。"""
    import shutil
    import workflow as wfmod
    src = _two_sheet_book(tmp_path)
    dest = Path(tmp_path) / "result.xlsx"
    shutil.copyfile(src, dest)
    with pytest.raises(store.Invalid) as ei:
        wfmod.write_back(src, dest, "S2", 1,
                         [(0, wfmod.OutputCol(1, "status"), "done")])
    assert "S2" in str(ei.value) and "S1" in str(ei.value)


def test_write_back_accepts_the_first_sheet(tmp_path):
    """配套围栏：传对表**不许**报错——否则上面那条会被「一律抛错」蒙混过去。"""
    import shutil
    import workflow as wfmod
    src = _two_sheet_book(tmp_path)
    dest = Path(tmp_path) / "result.xlsx"
    shutil.copyfile(src, dest)
    wfmod.write_back(src, dest, "S1", 1,
                     [(0, wfmod.OutputCol(1, "status"), "done")])
    from openpyxl import load_workbook
    assert load_workbook(dest)["S1"].cell(row=2, column=2).value == "done"


def test_archive_uses_the_sheet_the_row_numbers_came_from(disp, tmp_path):
    """工作流上存的 sheet 若是别的表名，归档**照样成功**——因为调度器传的是
    ``view.sheet``（第一张表），不是用户填的那个。"""
    uid = _user()
    wf = _workflow(uid, tmp_path)
    store.update_workflow(wf["id"], sheet="并不存在的表")
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    assert _archived(r0["id"]) == 1
    from openpyxl import load_workbook
    result = Path(store.get_workflow(wf["id"])["result_path"])
    assert load_workbook(result)["S1"].cell(row=2, column=4).value == "done"


def test_result_path_defaults_next_to_the_source(disp, tmp_path):
    """没配 result_path 时落到 source.xlsx 旁边，而不是变成一个叫 "." 的路径。"""
    uid = _user()
    wf = _workflow(uid, tmp_path)
    disp._workflow_run_pass()
    r0 = _first_row(wf)
    store.update_task(r0["id"], status="done", summary="ok")
    disp._workflow_archive_pass()
    assert disp._archive_wait(30)
    result = store.get_workflow(wf["id"])["result_path"]
    assert result and Path(result).name == "result.xlsx"
    assert Path(result).exists()
