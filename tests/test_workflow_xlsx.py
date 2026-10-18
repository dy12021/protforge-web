"""读表 / 取值 / 渲染，全部是纯函数，不碰数据库（规格 4.5 / 5.1 / 5.2 / 5.3）。

本任务最容易写错的三处，各有一条专门的用例：

- **公式必须回退**：openpyxl 自己写的公式，用 ``data_only=True`` 读回来是 None。
  不回退的表现是「整列提示词全是 None，而任务照跑不误，只是算了个空」。
- **全空数据行要剔除**：真实 Excel 尾部常有一两行格式残留，``ws.max_row``
  会把它们算成数据行——不剔的话工作流会跑出一堆「输入全空」的行任务。
- **context 列必须整列恒定**：不恒定就报错，不能悄悄取第一行。
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import store
import workflow


def _wb(tmp_path, rows, name="t.xlsx"):
    p = Path(tmp_path) / name
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in rows:
        ws.append(list(r))
    wb.save(p)
    return p


# --- 上传校验（规格 4.5 ②）-------------------------------------------------

def test_check_upload_accepts_real_xlsx(tmp_path):
    p = _wb(tmp_path, [["编号", "序列"], ["1", "AAAA"]])
    workflow.check_upload("表.xlsx", p.read_bytes())      # 不抛即通过


def test_check_upload_rejects_empty():
    with pytest.raises(store.Invalid):
        workflow.check_upload("空.xlsx", b"")


def test_check_upload_rejects_shell_by_extension():
    with pytest.raises(store.Invalid):
        workflow.check_upload("evil.sh", b"echo hi")


def test_check_upload_rejects_wrong_magic():
    """扩展名对但内容不是 zip——只看扩展名等于没校验。"""
    with pytest.raises(store.Invalid):
        workflow.check_upload("假.xlsx", b"not a zip at all")


def test_check_upload_rejects_oversize():
    with pytest.raises(store.Invalid):
        workflow.check_upload("大.xlsx", b"PK\x03\x04" + b"0" * (workflow.UPLOAD_MAX + 1))


def test_safe_display_name_strips_path():
    """落盘只用固定名 source.xlsx，但展示名要去掉路径分隔符。"""
    assert "/" not in workflow.safe_display_name("../../evil.xlsx")
    assert "\\" not in workflow.safe_display_name("..\\..\\evil.xlsx")
    assert len(workflow.safe_display_name("x" * 500 + ".xlsx")) <= 120


# --- 读表 -------------------------------------------------------------------

def test_read_source_basics(tmp_path):
    p = _wb(tmp_path, [["编号", "序列"], ["1", "AAA"], ["2", "BBB"]])
    v = workflow.read_source(p)
    assert v.sheet == "S1"
    assert v.headers == ["编号", "序列"]
    assert v.col_count == 2
    assert len(v.row_numbers) == 2


def test_read_source_drops_trailing_blank_rows(tmp_path):
    p = Path(tmp_path) / "t.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "序列"])
    ws.append(["1", "AAA"])
    ws.append([None, None])      # 真实表格尾部常有这种格式残留
    wb.save(p)
    v = workflow.read_source(p)
    assert len(v.row_numbers) == 1


def test_read_source_header_row_is_configurable(tmp_path):
    p = _wb(tmp_path, [["说明行"], ["编号", "序列"], ["1", "AAA"]])
    v = workflow.read_source(p, header_row=2)
    assert v.headers == ["编号", "序列"]
    assert len(v.row_numbers) == 1


def test_read_source_truncates_long_cells(tmp_path):
    long = "x" * (workflow.CELL_MAX + 500)
    p = _wb(tmp_path, [["编号", "备注"], ["1", long]])
    v = workflow.read_source(p)
    assert len(v.preview[0][1]) == workflow.CELL_MAX


def test_read_source_rejects_empty_sheet(tmp_path):
    p = Path(tmp_path) / "empty.xlsx"
    Workbook().save(p)
    with pytest.raises(store.Invalid):
        workflow.read_source(p)


# --- 公式回退（规格 5.3）----------------------------------------------------

def test_value_falls_back_to_formula_text(tmp_path):
    p = Path(tmp_path) / "f.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "长度"])
    ws.append(["1", "=LEN(B1)"])       # openpyxl 自己写的公式，没有缓存值
    wb.save(p)
    v = workflow.read_source(p)
    got = workflow.row_values(p, v, 0)
    assert got[1] == "=LEN(B1)", "回退失效：提示词里会喂进 None"
    assert got[1] is not None


def test_value_prefers_cached_result(tmp_path):
    p = _wb(tmp_path, [["编号", "长度"], ["1", "42"]])
    v = workflow.read_source(p)
    assert workflow.row_values(p, v, 0)[1] == "42"


# --- context 列恒定性（规格 5.1）--------------------------------------------

def test_context_column_must_be_constant(tmp_path):
    p = _wb(tmp_path, [["编号", "项目", "序列"],
                       ["1", "XBR", "AAA"], ["2", "XBR", "BBB"]])
    v = workflow.read_source(p)
    workflow.check_context_columns(p, v, [1])         # 恒定，不抛
    with pytest.raises(store.Invalid):
        workflow.check_context_columns(p, v, [2])     # 序列在变，当 context 要报错


def test_context_check_ignores_blank_cells(tmp_path):
    """空单元格不算「变了」——否则整列只有一个值、其余空着反而报错。"""
    p = _wb(tmp_path, [["编号", "项目"], ["1", "XBR"], ["2", None]])
    v = workflow.read_source(p)
    workflow.check_context_columns(p, v, [1])


# --- 渲染 -------------------------------------------------------------------

def test_render_substitutes_slots():
    out = workflow.render("查 {{seq}}，项目 {{proj}}",
                          {"seq": "AAA", "proj": "XBR"})
    assert out == "查 AAA，项目 XBR"


def test_render_allows_repeated_slot():
    out = workflow.render("{{seq}} / {{seq}}", {"seq": "A"})
    assert out == "A / A"


def test_render_missing_slot_becomes_blank_not_error():
    """空单元格渲染成空串、不报错——空表是常态，不是用户的错。"""
    assert workflow.render("查 [{{seq}}]", {}) == "查 []"


def test_extract_slots_dedupes_and_strips():
    got = workflow.extract_slots("{{a}} {{ b }} {{a}}")
    assert got == ["a", "b"]


def test_validate_binding_rejects_unbound_slot():
    cols = [{"col_index": 0, "role": "input", "slot": "seq"}]
    with pytest.raises(store.Invalid):
        workflow.validate_binding("查 {{seq}} 与 {{proj}}", cols)


def test_validate_binding_rejects_unused_column():
    cols = [{"col_index": 0, "role": "input", "slot": "seq"},
            {"col_index": 1, "role": "context", "slot": "没人用"}]
    with pytest.raises(store.Invalid):
        workflow.validate_binding("查 {{seq}}", cols)


def test_validate_binding_accepts_healthy_binding():
    cols = [{"col_index": 0, "role": "key"},
            {"col_index": 1, "role": "input", "slot": "seq"},
            {"col_index": 2, "role": "context", "slot": "proj"},
            {"col_index": 3, "role": "output", "field": "summary"}]
    workflow.validate_binding("查 {{seq}}，项目 {{proj}}", cols)


def test_output_column_needs_a_known_field():
    cols = [{"col_index": 0, "role": "output", "field": "随便写的"}]
    with pytest.raises(store.Invalid):
        workflow.validate_binding("x", cols)
