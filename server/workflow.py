"""工作流的纯逻辑：读 xlsx、判定列角色、渲染提示词、写回结果表。

**这里一个数据库调用都没有**，全是可以单独测的纯函数——因为这是最容易
悄悄算错的地方：公式读成 None、全空行混进数据行、context 列偷偷取第一行，
三种都不会报错，只是结果不对。
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

import store

UPLOAD_MAX = 20 * 1024 * 1024
CELL_MAX = 8000
PREVIEW_ROWS = 50
ROLES = ("key", "input", "context", "output", "ignore")
OUTPUT_FIELDS = ("status", "summary", "error", "artifacts")
XLSX_MAGIC = b"PK\x03\x04"
_SLOT_RE = re.compile(r"\{\{\s*([^{}\s]+)\s*\}\}")


@dataclass
class SheetView:
    sheet: str
    headers: list[str]
    row_numbers: list[int]      # 0 基数据行序号 -> sheet 行号（跳过全空行）
    preview: list[list[str]]
    col_count: int


# --- 上传校验（规格 4.5）----------------------------------------------------

def safe_display_name(filename: str) -> str:
    """**只用于展示**。落盘一律用固定名 source.xlsx（规格 4.5 ①）。

    file.filename 是用户可控的，直接拼进落盘路径就是目录穿越。
    """
    name = (filename or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f<>:\"|?*]", "", name).strip()
    return name[:120] or "table.xlsx"


def check_upload(filename: str, data: bytes) -> None:
    """框架对上传内容零校验：空文件和 .sh 实测都是 200 照收。自己判。"""
    if not data:
        raise store.Invalid("文件是空的")
    if len(data) > UPLOAD_MAX:
        raise store.Invalid(f"文件超过 {UPLOAD_MAX // (1024 * 1024)} MB 上限")
    ext = Path(safe_display_name(filename)).suffix.lower()
    if ext != ".xlsx":
        raise store.Invalid("只支持 .xlsx")
    if not data.startswith(XLSX_MAGIC):
        raise store.Invalid("这个文件不是 xlsx（内容与扩展名不符）")


# --- 读表 -------------------------------------------------------------------

def _cell_text(val, raw) -> str:
    """单元格取文本：**公式回退**（规格 5.3）+ `CELL_MAX` 截断。

    openpyxl 自己写的公式没有缓存值，`data_only=True` 读回来是 None——
    不回退的话整列提示词全是 None，而任务照跑不误，只是算了个空。

    read_source / row_values / check_context_columns **三处共用这一份**，
    免得改了一处另外两处悄悄不同。
    """
    if val is None and isinstance(raw, str) and raw.startswith("="):
        val = raw
    return str(val)[:CELL_MAX] if val is not None else ""


def _row_texts(vals, raws, max_col: int) -> list[str]:
    """一整行前 max_col 个单元格的文本。

    zip_longest 给两个工作簿补的 None、iter_rows 给短行补的空位，都在这里
    兜成空串，所以调用方拿到的永远是 max_col 个 str。
    """
    vals = vals or ()
    raws = raws or ()
    return [_cell_text(vals[i] if i < len(vals) else None,
                       raws[i] if i < len(raws) else None)
            for i in range(max_col)]


def _header_text(text: str, col: int) -> str:
    """列名：去掉首尾空白，空的（None / 空串 / 纯空白）占位成「列N」。

    纯空白表头塌成空串会往前端塞一个**空列名**——前端列表要拿它当列名显示，
    用户看到的是一列表头里有一列是空的，还以为程序没读到这一列。
    """
    return text.strip() or f"列{col}"


def _row_at(ws, r: int, max_col: int) -> list:
    """读 sheet 第 r 行的前 max_col 个值。**只给 row_values 用**（逐行取值）。

    走 iter_rows 而不是逐格 cell()：read_only 工作簿上逐格调用极慢，
    而一次一行的迭代只解析一遍共享字符串表。

    **不要写 ws.cell(..., values_only=True)**——cell() 返回 Cell 对象，
    values_only 是 iter_rows 的参数，传给它直接 TypeError。
    """
    for tup in ws.iter_rows(min_row=r, max_row=r, max_col=max_col,
                            values_only=True):
        return list(tup)
    return []


def _blank(vals) -> bool:
    return all((v or "") == "" for v in vals)


def read_source(path: Path, header_row: int = 1) -> SheetView:
    """读表头与数据行。**全空数据行在这里就剔掉**（规格 5.1）。

    真实 Excel 尾部常有一两行只有格式没内容的残留，``ws.max_row`` 会把它们
    算成数据行——不剔的话工作流会跑出一堆「输入全空」的行任务，每行都白烧
    一次完整计算。
    """
    wb_vals = load_workbook(path, data_only=True, read_only=True)
    try:
        wb_raw = load_workbook(path, data_only=False, read_only=True)
    except Exception:
        wb_raw = None
    try:
        wsv = wb_vals[wb_vals.sheetnames[0]]
        wsr = wb_raw[wb_raw.sheetnames[0]] if wb_raw else wsv
        sheet = wsv.title
        if header_row < 1:
            raise store.Invalid("表头行号从 1 开始")
        # 缺 <dimension> 时 openpyxl 给不出列数（None），定位也做不了。
        # 上一版崩在「按列数循环读表头」那行，现在那里没有了，改在这里显式抛，
        # **崩溃条件保持不变**（只有 max_column 为 None 才崩），不静默按 1 列读。
        if wsv.max_column is None:
            raise TypeError("这张表没有 <dimension> 元素，openpyxl 读不出列数")

        max_col = wsv.max_column or 1
        # **单遍顺序读完，表头也在同一次迭代里，绝不按行号或列号反复定位。**
        # read_only 工作簿上每定位一行（或一格）都要重开 archive 流、新建 parser
        # 从 XML 流头重扫：iter_rows(min_row=r) 逐行定位是 O(n^2)
        # （实测 800 行 19.8 秒），逐格 cell() 更狠（实测 800 行 253.85 秒）——
        # 而 Task 5 每放行一行就调一次 read_source，用户看到的是
        # 「保存转圈很久没有任何提示」。表头单独再读一遍虽然只有几毫秒，
        # 却是同一类退化的种子，所以并进这一次迭代。
        headers: list[str] = [f"列{i + 1}" for i in range(max_col)]
        row_numbers: list[int] = []
        preview: list[list[str]] = []
        it_vals = wsv.iter_rows(min_row=header_row, max_row=wsv.max_row,
                                max_col=max_col, values_only=True)
        it_raw = wsr.iter_rows(min_row=header_row, max_row=wsv.max_row,
                               max_col=max_col, values_only=True)
        for r, pair in enumerate(zip_longest(it_vals, it_raw), start=header_row):
            # zip_longest 补 None：两个工作簿行数不一致时短的那个按空行算。
            vals, raws = pair
            texts = _row_texts(vals, raws, max_col)
            if r == header_row:
                headers = [_header_text(t, i + 1) for i, t in enumerate(texts)]
                continue
            if _blank(texts):
                continue
            row_numbers.append(r)
            if len(preview) < PREVIEW_ROWS:
                preview.append(texts)

        if not row_numbers:
            raise store.Invalid("表里没有有效数据行")
        return SheetView(sheet=sheet, headers=headers, row_numbers=row_numbers,
                         preview=preview, col_count=len(headers))
    finally:
        wb_vals.close()
        if wb_raw is not None and wb_raw is not wb_vals:
            wb_raw.close()


def row_values(path: Path, view: SheetView, row_index: int) -> dict[int, str]:
    """取某一行的 {0 基列号: 文本}。"""
    if row_index < 0 or row_index >= len(view.row_numbers):
        raise store.Invalid(f"行号越界：{row_index}")
    wb = load_workbook(path, data_only=True, read_only=True)
    wbr = load_workbook(path, data_only=False, read_only=True)
    try:
        r = view.row_numbers[row_index]
        vals = _row_at(wb[view.sheet], r, view.col_count)
        raws = _row_at(wbr[view.sheet], r, view.col_count)
        return dict(enumerate(_row_texts(vals, raws, view.col_count)))
    finally:
        wb.close()
        wbr.close()


def check_context_columns(path: Path, view: SheetView,
                          context_cols: list[int]) -> None:
    """context 列必须整列恒定，否则**报错**而不是悄悄取第一行（规格 5.1）。

    悄悄取第一行的后果是：用户以为整表共用一个值，实际前几行悄悄不同，
    而整批结果看起来都正常。
    """
    if not context_cols:
        return
    # **两个工作簿只开一次，整表单遍读完再比。** 逐行调 row_values 等于
    # 每行重新 load 一遍——1000 行的表就是 2000 次整表加载，20MB 的表
    # 根本跑不完，而且用户看到的是「保存转圈很久没有任何提示」。
    # 逐行调 _row_at 也一样糟：read_only 下定位第 r 行要从流头重扫，是 O(n^2)。
    # 所以先单遍把要比较的值收齐，再逐列比——和原来的「先列后行」嵌套顺序
    # 一致，报错时挑中的列与行号才不变。
    if not view.row_numbers:
        return
    order: dict[int, int] = {}
    for idx, r in enumerate(view.row_numbers):
        if r not in order:      # 重复行号取第一次，与逐行比较的下标一致
            order[r] = idx
    first = min(order)
    last = max(order)
    # 按数据行下标存，row_numbers 里的下标原样保留（跳过的全空行不参与）。
    col_values: dict[int, list[str]] = {
        c: [""] * len(view.row_numbers) for c in context_cols}
    wb = load_workbook(path, data_only=True, read_only=True)
    wbr = load_workbook(path, data_only=False, read_only=True)
    try:
        it_vals = wb[view.sheet].iter_rows(min_row=first, max_row=last,
                                           max_col=view.col_count,
                                           values_only=True)
        it_raw = wbr[view.sheet].iter_rows(min_row=first, max_row=last,
                                          max_col=view.col_count,
                                          values_only=True)
        for r, pair in enumerate(zip_longest(it_vals, it_raw), start=first):
            idx = order.get(r)
            if idx is None:
                continue
            vals, raws = pair
            vals = vals or ()
            raws = raws or ()
            for col in context_cols:
                col_values[col][idx] = _cell_text(
                    vals[col] if col < len(vals) else None,
                    raws[col] if col < len(raws) else None)
        for col in context_cols:
            seen = None
            seen_row = 0
            for idx, val in enumerate(col_values[col]):
                if val == "":
                    continue
                if seen is None:
                    seen, seen_row = val, idx
                elif val != seen:
                    raise store.Invalid(
                        f"第 {col + 1} 列「{view.headers[col]}」被设成上下文，"
                        f"但它每一行都不一样（第 {idx + 1} 行是「{val[:30]}」，"
                        f"第 {seen_row + 1} 行是「{seen[:30]}」）。"
                        f"要么改成输入列，要么删掉这列。")
    finally:
        wb.close()
        wbr.close()


# --- 渲染 -------------------------------------------------------------------

def extract_slots(template: str) -> list[str]:
    out: list[str] = []
    for m in _SLOT_RE.finditer(template or ""):
        name = m.group(1)
        if name not in out:
            out.append(name)
    return out


def render(template: str, values: dict[str, str]) -> str:
    """未绑定的槽位渲染成空串，不报错——空单元格是常态不是用户的错。"""
    def sub(m):
        return values.get(m.group(1), "")
    return _SLOT_RE.sub(sub, template or "")


def validate_binding(template: str, cols: list[dict]) -> None:
    slots = set(extract_slots(template))
    bound = {c.get("slot", "") for c in cols
             if c.get("role") in ("input", "context") and c.get("slot")}
    missing = slots - bound
    if missing:
        raise store.Invalid(
            "提示词里的这些变量没有列绑定：" + "、".join(f"{{{{{m}}}}}" for m in sorted(missing)))
    unused = bound - slots
    if unused:
        raise store.Invalid(
            "这些列绑了变量但提示词里没用到：" + "、".join(sorted(unused))
            + "。要不要改成忽略？")
    keys = [c for c in cols if c.get("role") == "key"]
    if len(keys) != 1:
        raise store.Invalid(f"必须有且只有一列是「标识列」，现在有 {len(keys)} 列")
    for c in cols:
        if c.get("role") == "output" and c.get("field") not in OUTPUT_FIELDS:
            raise store.Invalid(
                f"结果列「{c.get('header')}」的字段名不认识：{c.get('field')}。"
                "可用：" + "、".join(OUTPUT_FIELDS))
